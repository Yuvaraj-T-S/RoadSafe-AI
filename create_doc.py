from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

document = Document()

# Title
title = document.add_heading('RoadSafe AI - Machine Learning Model Documentation', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Overview
document.add_heading('1. Overview', level=1)
document.add_paragraph(
    'The RoadSafe AI predictive system utilizes a Machine Learning model to calculate and '
    'predict road safety risk scores in real-time. This predictive capability allows '
    'city planners and traffic management to perform scenario-based assessments by varying environmental '
    'and behavioral factors to see their immediate impact on overall road safety.'
)

# Model Architecture
document.add_heading('2. Model Architecture and Technology', level=1)
document.add_paragraph('Algorithm: ', style='List Bullet').add_run('Random Forest Regressor').bold = True
document.add_paragraph('Library: ', style='List Bullet').add_run('scikit-learn (Python)').bold = True
document.add_paragraph(
    'The Random Forest Regressor was chosen due to its robustness against overfitting '
    'and its strong ability to handle categorical variables (like weather conditions and lighting) '
    'once properly encoded. The model builds an ensemble of decision trees during training '
    'and outputs the mean prediction of individual trees.'
)

# Dataset & Features
document.add_heading('3. Dataset & Features', level=1)
document.add_paragraph(
    'The model is trained on historical accident data. The target variable is the '
    'Risk Score (ranging from 0 to 100), calculated fundamentally through the severity '
    'of past accidents (weighing fatalities and injuries heaviest).'
)

document.add_heading('Input Features:', level=2)
p1 = document.add_paragraph('', style='List Bullet')
p1.add_run('Road Segment: ').bold = True
p1.add_run('The name/location of the road.')
p2 = document.add_paragraph('', style='List Bullet')
p2.add_run('Weather Condition: ').bold = True
p2.add_run('Clear, Rain, Fog, Storm, etc.')
p3 = document.add_paragraph('', style='List Bullet')
p3.add_run('Lighting Condition: ').bold = True
p3.add_run('Adequate, Poor.')
p4 = document.add_paragraph('', style='List Bullet')
p4.add_run('Road Condition: ').bold = True
p4.add_run('Dry, Wet, Damaged, Under construction.')
p5 = document.add_paragraph('', style='List Bullet')
p5.add_run('Contributing Factor: ').bold = True
p5.add_run('Overspeeding, Distracted Driving, Signal Violations, etc.')

# Pipeline
document.add_heading('4. Preprocessing Pipeline', level=1)
document.add_paragraph(
    'Before data is fed into the Random Forest model, it passes through a scikit-learn Pipeline '
    'utilizing a ColumnTransformer. Since all input features are text-based categories, '
    'a One-Hot Encoder converts these string values into binary matrix representations '
    'that the mathematical regression algorithm can understand. The pipeline ensures '
    'that raw JSON data directly from the frontend is transformed on-the-fly perfectly.'
)

# Deployment
document.add_heading('5. Deployment and API Integration', level=1)
document.add_paragraph(
    'The model is serialized and saved as a compressed Joblib file (roadsafe_model.joblib). '
    'It is served in real-time by a FastAPI backend (main.py). When a user interacts '
    'with the React frontend and clicks "Predict Scenario", an asynchronous POST request '
    'is made to the backend. The API decodes the JSON payload, loads it into a Pandas DataFrame, '
    'passes it through the inference pipeline, and immediately returns the newly predicted risk score.'
)

document.save('RoadSafe_AI_Model_Documentation.docx')
print("Document generated successfully.")
